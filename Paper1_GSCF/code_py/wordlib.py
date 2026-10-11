"""wordlib.py: python-docx helpers. Built-in Heading 1-3 styles for every heading, a three-level TOC field with cached
entries and page numbers (two-pass via LibreOffice PDF conversion), booktabs-style tables, figures, citation rendering."""
import os, re, subprocess, shutil, tempfile
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from refs import REFS

FONT = 'Times New Roman'

def _set_font(style, name=FONT, size=None, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    f = style.font; f.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rfonts.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rfonts.get(qn(a)) is not None: del rfonts.attrib[qn(a)]
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color

def new_doc(body_size=12, line_numbers=False, margins=1.0):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, m, Inches(margins))
    st = doc.styles
    _set_font(st['Normal'], size=body_size)
    st['Normal'].paragraph_format.space_after = Pt(0)
    _set_font(st['Title'], size=16, bold=True)
    st['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st['Title'].paragraph_format.space_after = Pt(12)
    # remove the template's bottom border on Title
    ppr = st['Title'].element.get_or_add_pPr()
    for b in ppr.findall(qn('w:pBdr')): ppr.remove(b)
    for name, size, bold, italic in (('Heading 1', 14, True, False), ('Heading 2', 12, True, False), ('Heading 3', 12, False, True)):
        s = st[name]; _set_font(s, size=size, bold=bold, italic=italic)
        s.paragraph_format.space_before = Pt(14 if name == 'Heading 1' else 10); s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    _set_font(st['Caption'], size=10, bold=True, italic=False)
    for lvl in (1, 2, 3):
        nm = f'TOC {lvl}'
        try: s = st[nm]
        except KeyError: s = st.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = st['Normal']; _set_font(s, size=10, bold=(lvl == 1))
        s.paragraph_format.left_indent = Inches(0.25 * (lvl - 1)); s.paragraph_format.space_after = Pt(0); s.paragraph_format.line_spacing = 1.0
    if line_numbers:
        ln = OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous')
        sec._sectPr.append(ln)
    add_page_number_footer(doc)
    doc._bm = 0; doc._heads = []
    return doc

def add_page_number_footer(doc):
    p = doc.sections[0].footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for t, txt in (('begin', None), (None, ' PAGE '), ('end', None)):
        if t:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t); r._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)

# ------------------------------------------------------------------ headings, bookmarks, TOC
def heading(doc, text, level, bookmark=True):
    p = doc.add_paragraph(style=f'Heading {level}')
    if bookmark:
        doc._bm += 1; name = f'_Toc{900000 + doc._bm}'
        s = OxmlElement('w:bookmarkStart'); s.set(qn('w:id'), str(doc._bm)); s.set(qn('w:name'), name)
        p._p.append(s); p.add_run(text)
        e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(doc._bm)); p._p.append(e)
        doc._heads.append((level, text, name))
    else:
        p.add_run(text)
    return p

def add_toc(doc, pages=None, title='Contents'):
    """Insert TOC (levels 1 to 3). Must be called after all headings exist, so the caller builds the body first,
    then moves the TOC paragraphs to the placeholder. Returns list of paragraph elements."""
    pages = pages or {}
    paras = []
    heads = doc._heads
    for i, (lvl, text, bm) in enumerate(heads):
        p = doc.add_paragraph(style=f'TOC {lvl}')
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if i == 0:
            r = p.add_run()
            b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' TOC \\o "1-3" \\h \\z \\u '; r._r.append(it)
            sp = OxmlElement('w:fldChar'); sp.set(qn('w:fldCharType'), 'separate'); r._r.append(sp)
        h = OxmlElement('w:hyperlink'); h.set(qn('w:anchor'), bm); h.set(qn('w:history'), '1')
        r1 = OxmlElement('w:r'); t1 = OxmlElement('w:t'); t1.text = text; t1.set(qn('xml:space'), 'preserve'); r1.append(t1); h.append(r1)
        rt = OxmlElement('w:r'); rt.append(OxmlElement('w:tab')); h.append(rt)
        r2 = OxmlElement('w:r'); t2 = OxmlElement('w:t'); t2.text = str(pages.get(bm, '')); r2.append(t2); h.append(r2)
        p._p.append(h)
        if i == len(heads) - 1:
            r = p.add_run(); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)
        paras.append(p)
    return paras

def insert_toc_at(doc, anchor_par, pages=None, title='Contents', page_break_after=True):
    """Create TOC paragraphs (plus a title paragraph that is NOT a heading) and move them directly after anchor_par."""
    tpar = doc.add_paragraph(); r = tpar.add_run(title); r.bold = True; r.font.size = Pt(14)
    toc = add_toc(doc, pages)
    els = [tpar._p] + [p._p for p in toc]
    if page_break_after:
        pb = doc.add_paragraph(); pb.add_run().add_break(WD_BREAK.PAGE); els.append(pb._p)
    ref = anchor_par._p
    for e in els:
        ref.addnext(e); ref = e

def pdf_pages(docx_path):
    out = tempfile.mkdtemp()
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', out, docx_path], check=True, capture_output=True, timeout=300)
    pdf = os.path.join(out, os.path.splitext(os.path.basename(docx_path))[0] + '.pdf')
    n = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout).group(1))
    pages = []
    for i in range(1, n + 1):
        t = subprocess.run(['pdftotext', '-f', str(i), '-l', str(i), '-layout', pdf, '-'], capture_output=True, text=True).stdout
        pages.append(re.sub(r'\s+', ' ', t))
    return pages, pdf

def compute_pages(docx_path, heads, body_marker):
    pages, pdf = pdf_pages(docx_path)
    start = next(i for i, t in enumerate(pages) if body_marker in t)
    res, cur = {}, start
    for lvl, text, bm in heads:
        key = re.sub(r'\s+', ' ', text)
        for i in range(cur, len(pages)):
            if key in pages[i]:
                res[bm] = i + 1; cur = i; break
        else:
            res[bm] = cur + 1
    return res

# ------------------------------------------------------------------ citations
def cite_text(text):
    def parenthetical(m):
        keys = m.group(1).split(';')
        items = sorted(((REFS[k][1], REFS[k][0]) for k in keys))
        return '(' + '; '.join(f'{lab}, {yr}' for yr, lab in items) + ')'
    def narrative(m):
        k = m.group(1); return f'{REFS[k][0]} ({REFS[k][1]})'
    text = re.sub(r'\[\[c:([^\]]+)\]\]', parenthetical, text)
    text = re.sub(r'\[\[n:([^\]]+)\]\]', narrative, text)
    return text

def para(doc, text, style='Normal', spacing=2.0, align=None, first_indent=0.0, italic=False, bold=False, size=None, highlight=False, space_after=0):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format; pf.line_spacing = spacing; pf.space_after = Pt(space_after)
    if first_indent: pf.first_line_indent = Inches(first_indent)
    if align == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if align == 'justify': p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); r.italic = italic; r.bold = bold
    if size: r.font.size = Pt(size)
    if highlight:
        from docx.enum.text import WD_COLOR_INDEX
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p

# ------------------------------------------------------------------ tables
def _border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge, val in kw.items():
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), val[0]); e.set(qn('w:sz'), str(val[1])); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000'); b.append(e)

def table(doc, header, rows, widths, font=9, align_first='left', bold_rows=(), rule_after=(), note=None, header_rows=1, merges=(), left_cols=(0,)):
    """Booktabs-style table: top rule, rule under header, bottom rule, no vertical lines. Cells may contain '\n'."""
    t = doc.add_table(rows=len(header) if header and isinstance(header[0], list) else 1, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    hdrs = header if (header and isinstance(header[0], list)) else [header]
    for j, wd in enumerate(widths): t.columns[j].width = Inches(wd)
    for hi, h in enumerate(hdrs):
        for j, txt in enumerate(h):
            c = t.rows[hi].cells[j]; c.width = Inches(widths[j]); c.text = ''
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt); r.bold = True; r.font.size = Pt(font)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for j, txt in enumerate(row):
            c = cells[j]; c.width = Inches(widths[j]); c.text = ''
            parts = str(txt).split('\n')
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in left_cols else WD_ALIGN_PARAGRAPH.CENTER
            for k, part in enumerate(parts):
                r = p.add_run(part); r.font.size = Pt(font); r.bold = (ri in bold_rows)
                if k < len(parts) - 1: r.add_break()
            p.paragraph_format.space_after = Pt(0)
    nh = len(hdrs)
    for (ri_, c0, c1) in merges:
        a = t.rows[ri_].cells[c0]; b = t.rows[ri_].cells[c1]; txt = a.text; m = a.merge(b)
        for p_ in m.paragraphs[1:]: p_._p.getparent().remove(p_._p)
    for c in t.rows[0].cells: _border(c, top=('single', 12))
    for c in t.rows[nh - 1].cells: _border(c, bottom=('single', 6))
    for c in t.rows[-1].cells: _border(c, bottom=('single', 12))
    for ri in rule_after:
        for c in t.rows[nh + ri].cells: _border(c, bottom=('single', 4))
    # repeat header row on page breaks
    for hi in range(nh):
        trPr = t.rows[hi]._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    if note:
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
        r = p.add_run(note); r.font.size = Pt(font - 0.5)
    return t

def figure(doc, path, width=6.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width))
    return p

def caption(doc, text, size=10):
    p = doc.add_paragraph(style='Caption'); r = p.add_run(text); r.font.size = Pt(size)
    p.paragraph_format.keep_with_next = True
    return p

def references(doc, keys, spacing=1.5):
    items = sorted((REFS[k][2] for k in keys), key=lambda s: re.sub(r'[^a-z]', '', s.lower()))
    for t in items:
        p = doc.add_paragraph(); pf = p.paragraph_format
        pf.left_indent = Inches(0.4); pf.first_line_indent = Inches(-0.4); pf.line_spacing = spacing; pf.space_after = Pt(3)
        p.add_run(t)

def landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
